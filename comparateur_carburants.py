"""Comparateur de prix des carburants autour d'une ville.

Ce script interroge en temps réel l'API open data du ministère de
l'Économie pour trouver les stations-service les moins chères dans un
rayon donné autour d'une commune française.

Sa particularité : il évalue la **fiabilité** de chaque prix. Les prix
étant déclarés par les stations elles-mêmes, certains ne sont plus mis
à jour depuis des semaines. Le script sépare donc les prix récents des
prix douteux avant d'établir le classement.

Étapes du traitement :
    1. Géocodage de la ville saisie (API de l'IGN), avec gestion des
       communes homonymes.
    2. Récupération paginée des stations situées dans le rayon choisi.
    3. Calcul vectorisé des distances (formule de haversine).
    4. Passage au format long (une ligne par station et par carburant)
       et calcul de l'ancienneté de chaque prix.
    5. Classement du top 3 par carburant, séparément pour les prix
       fiables et les prix douteux.
    6. Affichage console et export d'un rapport Excel mis en forme.

Exemple:
    $ python comparateur_carburants.py
    Ville : Aubergenville
    Rayon en km (1 à 30) : 10

Dépendances:
    requests, pandas, numpy, openpyxl

Sources des données:
    - Prix des carburants en France, flux instantané (v2),
      data.economie.gouv.fr
    - Service de géocodage de la Géoplateforme (IGN), data.geopf.fr

Auteur:
    Anis NEFIKHA, étudiant en BUT Science des Données (IUT de Lille).
"""

from __future__ import annotations

import os
import subprocess
import sys
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from typing import Any

import numpy as np
import pandas as pd
import requests
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# ---------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------

URL_CARBURANTS = (
    "https://data.economie.gouv.fr/api/explore/v2.1/catalog/datasets/"
    "prix-des-carburants-en-france-flux-instantane-v2/records"
)
URL_GEOCODAGE = "https://data.geopf.fr/geocodage/search"

# Nom du champ de prix dans l'API -> libellé affiché.
# Le champ de date associé se déduit en remplaçant "_prix" par "_maj".
CARBURANTS: dict[str, str] = {
    "gazole_prix": "Diesel (Gazole)",
    "sp95_prix": "Sans plomb 95 (SP95)",
    "e10_prix": "Sans plomb 95-E10 (E10)",
    "sp98_prix": "Sans plomb 98 (SP98)",
    "e85_prix": "Superéthanol (E85)",
}

RAYON_MAX_KM = 30
SEUIL_FIABILITE_JOURS = 10        # en dessous : prix considéré fiable
SEUIL_NON_REPRESENTATIF_JOURS = 20  # au-dessus : prix probablement faux
TIMEOUT_SECONDES = 10
TAILLE_PAGE_API = 100             # nombre max de résultats par appel
NB_PROPOSITIONS_VILLE = 5
RAYON_TERRE_KM = 6371
EXPORTER_EXCEL = True

# Libellés des colonnes du rapport
COL_RANG = "CLASSEMENT"
COL_STATION = "STATION"
COL_CARBURANT = "CARBURANT"
COL_PRIX = "PRIX AU LITRE (€)"
COL_DISTANCE = "DISTANCE (KM)"
COL_MAJ = "MISE À JOUR"
COL_STATUT = "FIABILITÉ"
COL_FIABLE = "fiable"

COLONNES_FIABLES = [
    COL_RANG, COL_STATION, COL_CARBURANT, COL_PRIX, COL_DISTANCE, COL_MAJ,
]
COLONNES_DOUTEUX = COLONNES_FIABLES + [COL_STATUT]
RANGS = ["1er", "2e", "3e"]

# Styles Excel
_TRAIT = Side(style="thin")
BORDURE = Border(left=_TRAIT, right=_TRAIT, top=_TRAIT, bottom=_TRAIT)
FOND_VERT = PatternFill("solid", fgColor="C6EFCE")
TEXTE_VERT = Font(bold=True, color="006100")
FOND_ORANGE = PatternFill("solid", fgColor="FFE5CC")
TEXTE_ORANGE = Font(bold=True, color="C65911")
FOND_ROUGE = PatternFill("solid", fgColor="FFC7CE")
TEXTE_ROUGE = Font(bold=True, color="9C0006")
TEXTE_GRIS = Font(italic=True, color="808080")
FOND_ENTETE = PatternFill("solid", fgColor="1F4E78")
TEXTE_ENTETE = Font(bold=True, color="FFFFFF")
TEXTE_TITRE = Font(bold=True, size=14, color="1F4E78")
TEXTE_SOUS_TITRE = Font(bold=True, size=12, color="1F4E78")


@dataclass(frozen=True)
class Lieu:
    """Commune de référence choisie par l'utilisateur.

    Attributes:
        nom: Libellé affiché, par exemple « Aubergenville (78410) ».
        longitude: Longitude en degrés décimaux (WGS84).
        latitude: Latitude en degrés décimaux (WGS84).
    """

    nom: str
    longitude: float
    latitude: float


# ---------------------------------------------------------------------
# Accès aux API
# ---------------------------------------------------------------------

def appeler_api(url: str, params: dict[str, Any]) -> dict[str, Any] | None:
    """Envoie une requête GET et renvoie la réponse JSON.

    Toutes les erreurs réseau (absence de connexion, délai dépassé,
    code HTTP d'erreur) sont interceptées : la fonction affiche un
    message lisible et renvoie None au lieu d'interrompre le programme.

    Args:
        url: Adresse de l'API.
        params: Paramètres ajoutés à l'URL.

    Returns:
        Le contenu JSON décodé, ou None en cas d'échec.
    """
    try:
        reponse = requests.get(url, params=params, timeout=TIMEOUT_SECONDES)
        reponse.raise_for_status()
        return reponse.json()
    except requests.RequestException as erreur:
        print(f"❌ Problème de connexion à l'API : {erreur}")
        return None


def recuperer_stations(lieu: Lieu, rayon_km: float) -> pd.DataFrame:
    """Récupère toutes les stations situées dans un rayon donné.

    L'API renvoie au plus 100 résultats par appel : les pages sont
    parcourues avec le paramètre ``offset``. L'API limitant
    ``offset + limit`` à 10 000, cette méthode convient à un rayon local
    mais pas à la France entière (utiliser alors l'endpoint /exports).

    Args:
        lieu: Point central de la recherche.
        rayon_km: Rayon de recherche en kilomètres.

    Returns:
        Un DataFrame avec une ligne par station (vide en cas d'erreur).
    """
    champs_prix = list(CARBURANTS)
    champs_maj = [champ.replace("_prix", "_maj") for champ in champs_prix]
    point = f"GEOM'POINT({lieu.longitude} {lieu.latitude})'"
    params: dict[str, Any] = {
        "where": f"within_distance(geom, {point}, {int(rayon_km * 1000)}m)",
        "select": ", ".join(
            ["adresse", "ville", "cp", "geom"] + champs_prix + champs_maj
        ),
        "limit": TAILLE_PAGE_API,
        "offset": 0,
    }

    stations: list[dict[str, Any]] = []
    while True:
        donnees = appeler_api(URL_CARBURANTS, params)
        if donnees is None:
            return pd.DataFrame()
        stations.extend(donnees["results"])
        params["offset"] += TAILLE_PAGE_API
        if params["offset"] >= donnees["total_count"]:
            break
    return pd.DataFrame(stations)


# ---------------------------------------------------------------------
# Saisies utilisateur
# ---------------------------------------------------------------------

def normaliser(texte: str) -> str:
    """Simplifie un nom de commune pour faciliter les comparaisons.

    Args:
        texte: Nom à normaliser.

    Returns:
        Le nom sans accents, en minuscules, avec les tirets remplacés
        par des espaces. Exemple : « Saint-Étienne » -> « saint etienne ».
    """
    decompose = unicodedata.normalize("NFKD", texte)
    sans_accents = "".join(
        car for car in decompose if not unicodedata.combining(car)
    )
    return sans_accents.lower().replace("-", " ").strip()


def choisir_commune(communes: list[dict[str, Any]]) -> dict[str, Any] | None:
    """Fait choisir une commune dans une liste numérotée.

    Args:
        communes: Communes candidates (clés ``nom``, ``cp``, ``contexte``).

    Returns:
        La commune choisie, ou None si l'utilisateur veut ressaisir.
    """
    print("Plusieurs communes correspondent :")
    for numero, commune in enumerate(communes, start=1):
        print(f"   {numero}. {commune['nom']} ({commune['cp']}) "
              f"— {commune['contexte']}")

    while True:
        choix = input(f"Ton choix (1 à {len(communes)}, "
                      "0 pour retaper la ville) : ").strip()
        if choix.isdigit() and 0 <= int(choix) <= len(communes):
            return None if choix == "0" else communes[int(choix) - 1]
        print("⚠️ Entre un des numéros proposés")


def demander_lieu() -> Lieu:
    """Demande une commune à l'utilisateur et la géocode.

    Si un seul résultat porte exactement le nom saisi, il est retenu
    directement. Sinon (homonymes ou faute de frappe), l'utilisateur
    choisit dans une liste.

    Returns:
        La commune retenue avec ses coordonnées GPS.
    """
    while True:
        saisie = input("Ville : ").strip()
        if len(saisie) < 3:
            print("⚠️ Tape au moins 3 lettres")
            continue

        reponse = appeler_api(URL_GEOCODAGE, {
            "q": saisie,
            "type": "municipality",
            "limit": NB_PROPOSITIONS_VILLE,
        })
        if reponse is None:
            continue

        communes = []
        for resultat in reponse.get("features", []):
            proprietes = resultat["properties"]
            longitude, latitude = resultat["geometry"]["coordinates"]
            communes.append({
                "nom": proprietes.get("name", proprietes.get("label")),
                "cp": proprietes.get("postcode", "?"),
                "contexte": proprietes.get("context", ""),
                "longitude": longitude,
                "latitude": latitude,
            })

        if not communes:
            print("❌ Ville introuvable, réessaie")
            continue

        exactes = [
            commune for commune in communes
            if normaliser(commune["nom"]) == normaliser(saisie)
        ]
        if len(exactes) == 1:
            choisie: dict[str, Any] | None = exactes[0]
        else:
            choisie = choisir_commune(exactes or communes)
        if choisie is None:
            continue

        lieu = Lieu(
            nom=f"{choisie['nom']} ({choisie['cp']})",
            longitude=choisie["longitude"],
            latitude=choisie["latitude"],
        )
        print(f"📍 Ville retenue : {lieu.nom}")
        return lieu


def demander_rayon() -> float:
    """Demande un rayon de recherche valide.

    Les virgules décimales sont acceptées (« 7,5 »).

    Returns:
        Un rayon strictement positif et inférieur ou égal à RAYON_MAX_KM.
    """
    while True:
        saisie = input(f"Rayon en km (1 à {RAYON_MAX_KM}) : ")
        try:
            rayon = float(saisie.replace(",", "."))
        except ValueError:
            print("⚠️ Entre un nombre, pas du texte")
            continue
        if 0 < rayon <= RAYON_MAX_KM:
            return rayon
        print(f"⚠️ Entre un nombre entre 1 et {RAYON_MAX_KM}")


# ---------------------------------------------------------------------
# Traitement des données
# ---------------------------------------------------------------------

def distance_km(
    lat1: float | pd.Series,
    lon1: float | pd.Series,
    lat2: float | pd.Series,
    lon2: float | pd.Series,
) -> Any:
    """Calcule la distance à vol d'oiseau entre deux points GPS.

    Utilise la formule de haversine. Grâce à numpy, la fonction accepte
    aussi bien des nombres que des colonnes entières (calcul vectorisé).

    Args:
        lat1: Latitude du premier point, en degrés.
        lon1: Longitude du premier point, en degrés.
        lat2: Latitude du second point, en degrés.
        lon2: Longitude du second point, en degrés.

    Returns:
        La distance en kilomètres (nombre ou Series selon les entrées).
    """
    phi1, phi2 = np.radians(lat1), np.radians(lat2)
    delta_phi = phi2 - phi1
    delta_lambda = np.radians(lon2) - np.radians(lon1)
    a = (np.sin(delta_phi / 2) ** 2
         + np.cos(phi1) * np.cos(phi2) * np.sin(delta_lambda / 2) ** 2)
    return 2 * RAYON_TERRE_KM * np.arcsin(np.sqrt(a))


def ajouter_distances(stations: pd.DataFrame, lieu: Lieu) -> pd.DataFrame:
    """Ajoute à chaque station sa distance au lieu de référence.

    Args:
        stations: Stations brutes renvoyées par l'API (colonne ``geom``).
        lieu: Point de référence.

    Returns:
        Une copie du DataFrame avec les colonnes ``lat``, ``lon`` et
        ``distance_km`` (NaN si la position est inconnue).
    """
    resultat = stations.copy()
    resultat["lat"] = resultat["geom"].str.get("lat")
    resultat["lon"] = resultat["geom"].str.get("lon")
    resultat["distance_km"] = distance_km(
        lieu.latitude, lieu.longitude, resultat["lat"], resultat["lon"]
    ).round(1)
    return resultat


def construire_tableau_long(stations: pd.DataFrame) -> pd.DataFrame:
    """Passe au format long et évalue la fiabilité de chaque prix.

    Format d'entrée (large) : une ligne par station, une colonne par
    carburant. Format de sortie (long) : une ligne par couple
    (station, carburant), ce qui simplifie filtres et classements.

    Args:
        stations: Stations avec leur distance (voir ajouter_distances).

    Returns:
        Un DataFrame avec la station, le carburant, le prix, la distance,
        la date de mise à jour, un statut lisible et un booléen
        ``fiable`` (True si le prix a moins de SEUIL_FIABILITE_JOURS).
    """
    champs_prix = list(CARBURANTS)
    champs_maj = [champ.replace("_prix", "_maj") for champ in champs_prix]

    donnees = stations.copy()
    donnees["id_station"] = range(len(donnees))
    donnees[COL_STATION] = donnees["adresse"] + ", " + donnees["ville"]

    # Deux passages au format long (prix et dates), puis jointure
    prix = donnees.melt(
        id_vars=["id_station", COL_STATION, "distance_km"],
        value_vars=champs_prix, var_name="champ", value_name=COL_PRIX,
    )
    prix["code"] = prix["champ"].str.replace("_prix", "")

    dates = donnees.melt(
        id_vars=["id_station"], value_vars=champs_maj,
        var_name="champ", value_name="maj",
    )
    dates["code"] = dates["champ"].str.replace("_maj", "")

    long = prix.merge(
        dates[["id_station", "code", "maj"]], on=["id_station", "code"]
    )
    long = long.dropna(subset=[COL_PRIX])
    long[COL_CARBURANT] = long["champ"].map(CARBURANTS)
    long = long.rename(columns={"distance_km": COL_DISTANCE})

    # Ancienneté de chaque prix, en jours
    dates_paris = (pd.to_datetime(long["maj"], utc=True)
                   .dt.tz_convert("Europe/Paris"))
    age = (pd.Timestamp.now(tz="Europe/Paris") - dates_paris).dt.days
    long[COL_MAJ] = dates_paris.dt.strftime("%d/%m/%Y %Hh%M").fillna("-")

    # np.select retient la PREMIÈRE condition vraie : l'ordre compte
    age_texte = age.astype("Int64").astype(str)
    conditions: list[Any] = [
        age.isna(),
        age > SEUIL_NON_REPRESENTATIF_JOURS,
        age >= SEUIL_FIABILITE_JOURS,
    ]
    statuts: list[Any] = [
        "❌ DATE INCONNUE",
        "❌ PAS REPRÉSENTATIF (il y a " + age_texte
        + " j, probablement faux)",
        "⚠️ ATTENTION (il y a " + age_texte + " j)",
    ]
    long[COL_STATUT] = np.select(conditions, statuts, default="✅ OK")
    long[COL_FIABLE] = age < SEUIL_FIABILITE_JOURS  # NaN < 10 -> False

    return long[[COL_STATION, COL_CARBURANT, COL_PRIX, COL_DISTANCE,
                 COL_MAJ, COL_STATUT, COL_FIABLE]]


def top3_par_carburant(
    selection: pd.DataFrame,
    tous_les_prix: pd.DataFrame,
    colonnes: list[str],
    avec_lignes_vides: bool,
) -> pd.DataFrame:
    """Établit le top 3 de chaque carburant.

    Le classement se fait par prix croissant puis, à prix égal, par
    distance croissante. La boucle porte sur les carburants (quelques
    itérations) et non sur les lignes, afin de conserver l'ordre de
    CARBURANTS et de signaler les carburants absents.

    Args:
        selection: Prix à classer (par exemple uniquement les fiables).
        tous_les_prix: Ensemble des prix, pour distinguer « non proposé »
            de « aucun prix fiable ».
        colonnes: Colonnes à conserver, dans l'ordre d'affichage.
        avec_lignes_vides: Si True, ajoute une ligne explicative pour les
            carburants sans résultat.

    Returns:
        Le classement, carburant par carburant.
    """
    blocs: list[pd.DataFrame] = []
    for carburant in CARBURANTS.values():
        prix_carburant = selection[selection[COL_CARBURANT] == carburant]

        if prix_carburant.empty:
            if avec_lignes_vides:
                vendu = (tous_les_prix[COL_CARBURANT] == carburant).any()
                message = ("Aucun prix fiable (voir tableau 2)" if vendu
                           else "Non proposé")
                blocs.append(pd.DataFrame([{
                    COL_RANG: "-", COL_STATION: message,
                    COL_CARBURANT: carburant, COL_PRIX: "-",
                    COL_DISTANCE: "-", COL_MAJ: "-",
                }]))
            continue

        top = (prix_carburant.sort_values([COL_PRIX, COL_DISTANCE])
               .head(len(RANGS)).copy())
        top.insert(0, COL_RANG, RANGS[:len(top)])
        blocs.append(top)

    if not blocs:
        return pd.DataFrame(columns=colonnes)
    return pd.concat(blocs, ignore_index=True)[colonnes]


# ---------------------------------------------------------------------
# Restitution : console
# ---------------------------------------------------------------------

def afficher_classement(classement: pd.DataFrame, titre: str) -> None:
    """Affiche un classement dans la console, groupé par carburant.

    Args:
        classement: Résultat de top3_par_carburant.
        titre: Titre affiché au-dessus du classement.
    """
    print(f"\n{titre}")
    if classement.empty:
        print("   Rien à signaler 🎉")
        return

    carburant_courant = None
    for ligne in classement.to_dict("records"):
        if ligne[COL_CARBURANT] != carburant_courant:
            carburant_courant = ligne[COL_CARBURANT]
            print(f"\n⛽ {carburant_courant}")
        if ligne[COL_RANG] == "-":
            print(f"   {ligne[COL_STATION]}")
            continue

        distance = ("? km" if pd.isna(ligne[COL_DISTANCE])
                    else f"{ligne[COL_DISTANCE]} km")
        statut = f" | {ligne[COL_STATUT]}" if COL_STATUT in ligne else ""
        print(f"   {ligne[COL_RANG]} | {ligne[COL_PRIX]:.3f} €/L | "
              f"{distance} | {ligne[COL_STATION]} | "
              f"MAJ {ligne[COL_MAJ]}{statut}")


# ---------------------------------------------------------------------
# Restitution : Excel
# ---------------------------------------------------------------------

def ouvrir_fichier(chemin: str) -> None:
    """Ouvre un fichier avec l'application par défaut du système.

    Fonctionne sous Windows, macOS et Linux. Un échec n'est pas
    bloquant : le chemin du fichier est déjà affiché.

    Args:
        chemin: Chemin absolu du fichier.
    """
    try:
        if sys.platform.startswith("win"):
            os.startfile(chemin)  # type: ignore[attr-defined]
        elif sys.platform == "darwin":
            subprocess.run(["open", chemin], check=False)
        else:
            subprocess.run(["xdg-open", chemin], check=False)
    except OSError:
        print("ℹ️ Ouvre le fichier manuellement (chemin ci-dessus).")


def ecrire_titre(feuille: Worksheet, ligne: int, texte: str,
                 police: Font) -> None:
    """Écrit un titre dans la colonne A d'une ligne.

    Args:
        feuille: Onglet Excel.
        ligne: Numéro de ligne (à partir de 1).
        texte: Texte du titre.
        police: Police à appliquer.
    """
    feuille.cell(row=ligne, column=1, value=texte).font = police
    feuille.row_dimensions[ligne].height = 24


def mettre_en_forme(feuille: Worksheet, ligne_entete: int,
                    nb_lignes: int) -> None:
    """Met en forme un tableau : en-tête, bordures, formats et couleurs.

    Règles de couleur :
        - tableau des prix douteux : rouge si non représentatif ou date
          inconnue, orange sinon ;
        - autres tableaux : la 1re place de chaque carburant en vert.

    Args:
        feuille: Onglet contenant le tableau.
        ligne_entete: Ligne de l'en-tête du tableau.
        nb_lignes: Nombre de lignes de données sous l'en-tête.
    """
    entetes = [str(c.value) for c in feuille[ligne_entete]
               if c.value is not None]
    nb_colonnes = len(entetes)

    for cellule in feuille[ligne_entete][:nb_colonnes]:
        cellule.font = TEXTE_ENTETE
        cellule.fill = FOND_ENTETE
        cellule.border = BORDURE
        cellule.alignment = Alignment(horizontal="center",
                                      vertical="center", wrap_text=True)
    feuille.row_dimensions[ligne_entete].height = 30

    lignes = feuille.iter_rows(min_row=ligne_entete + 1,
                               max_row=ligne_entete + nb_lignes,
                               max_col=nb_colonnes)
    for ligne in lignes:
        for cellule in ligne:
            entete = entetes[cellule.column - 1]
            cellule.border = BORDURE
            cellule.alignment = Alignment(vertical="center")
            valeur = cellule.value
            if isinstance(valeur, float):
                cellule.number_format = ('0.0 "km"' if "km" in entete.lower()
                                         else '0.000 "€"')
            elif isinstance(valeur, str) and (
                    valeur in ("Non proposé", "-")
                    or valeur.startswith("Aucun")):
                cellule.font = TEXTE_GRIS

        if COL_STATUT in entetes:
            statut = str(ligne[entetes.index(COL_STATUT)].value)
            fond, police = ((FOND_ROUGE, TEXTE_ROUGE)
                            if statut.startswith("❌")
                            else (FOND_ORANGE, TEXTE_ORANGE))
        elif (COL_RANG in entetes
              and ligne[entetes.index(COL_RANG)].value == RANGS[0]):
            fond, police = FOND_VERT, TEXTE_VERT
        else:
            continue
        for cellule in ligne:
            cellule.fill = fond
            cellule.font = police


def ajuster_largeurs(feuille: Worksheet, ligne_depart: int) -> None:
    """Adapte la largeur des colonnes à leur contenu.

    Les titres (police de taille 12 ou plus) sont ignorés, et une marge
    est prévue pour le texte en gras.

    Args:
        feuille: Onglet à ajuster.
        ligne_depart: Première ligne prise en compte.
    """
    for colonne in feuille.iter_cols(min_row=ligne_depart):
        longueurs = [len(str(c.value)) for c in colonne
                     if c.value is not None and (c.font.size or 11) < 12]
        if longueurs:
            lettre = get_column_letter(colonne[0].column)
            feuille.column_dimensions[lettre].width = min(
                max(longueurs) * 1.2 + 3, 70)


def exporter_excel(
    fiables: pd.DataFrame,
    douteux: pd.DataFrame,
    stations: pd.DataFrame,
    lieu: Lieu,
    rayon_km: float,
) -> str | None:
    """Crée le rapport Excel et l'ouvre.

    Le classeur contient deux onglets :
        - « Top 3 » : prix fiables puis prix douteux ;
        - « Toutes les stations » : annexe triée par distance.

    Args:
        fiables: Classement des prix fiables.
        douteux: Classement des prix douteux.
        stations: Toutes les stations du rayon.
        lieu: Commune de référence.
        rayon_km: Rayon de recherche.

    Returns:
        Le chemin du fichier créé, ou None s'il n'a pas pu être écrit
        (par exemple parce qu'il est déjà ouvert dans Excel).
    """
    maintenant = datetime.now()
    titre = (f"PRIX DES CARBURANTS AU {maintenant:%d/%m/%Y À %HH%M} "
             f"— {lieu.nom}, rayon {rayon_km:g} km")
    nom_lieu = (lieu.nom.replace(" ", "_")
                .replace("(", "").replace(")", ""))
    fichier = (f"carburants_{nom_lieu}_{rayon_km:g}km_"
               f"{maintenant:%Y-%m-%d}.xlsx")

    libelles = {"adresse": "ADRESSE", "ville": "VILLE",
                "cp": "CODE POSTAL", "distance_km": COL_DISTANCE}
    libelles.update({champ: f"{nom} (€/L)"
                     for champ, nom in CARBURANTS.items()})
    annexe = (stations[["adresse", "ville", "cp", "distance_km"]
                       + list(CARBURANTS)]
              .sort_values("distance_km")
              .rename(columns=libelles)
              .fillna("Non proposé"))

    try:
        with pd.ExcelWriter(fichier, engine="openpyxl") as writer:
            # Onglet 1 : prix fiables puis prix douteux
            onglet = "Top 3"
            entete_fiables = 4
            fiables.to_excel(writer, sheet_name=onglet, index=False,
                             startrow=entete_fiables - 1)
            feuille = writer.sheets[onglet]
            ecrire_titre(feuille, 1, titre, TEXTE_TITRE)
            ecrire_titre(feuille, 3,
                         "✅ PRIX FIABLES (mis à jour il y a moins de "
                         f"{SEUIL_FIABILITE_JOURS} jours)", TEXTE_SOUS_TITRE)
            mettre_en_forme(feuille, entete_fiables, len(fiables))

            titre_douteux = entete_fiables + len(fiables) + 3
            if douteux.empty:
                ecrire_titre(feuille, titre_douteux,
                             "✅ Aucun prix douteux dans cette zone",
                             TEXTE_SOUS_TITRE)
            else:
                ecrire_titre(feuille, titre_douteux,
                             "⚠️ PRIX DOUTEUX (pas mis à jour depuis "
                             f"{SEUIL_FIABILITE_JOURS} jours ou plus)",
                             TEXTE_SOUS_TITRE)
                douteux.to_excel(writer, sheet_name=onglet, index=False,
                                 startrow=titre_douteux)
                mettre_en_forme(feuille, titre_douteux + 1, len(douteux))
            ajuster_largeurs(feuille, entete_fiables)

            # Onglet 2 : annexe complète
            annexe.to_excel(writer, sheet_name="Toutes les stations",
                            index=False, startrow=2)
            feuille_annexe = writer.sheets["Toutes les stations"]
            ecrire_titre(feuille_annexe, 1, titre, TEXTE_TITRE)
            mettre_en_forme(feuille_annexe, 3, len(annexe))
            ajuster_largeurs(feuille_annexe, 3)
            feuille_annexe.freeze_panes = "A4"
            derniere_colonne = get_column_letter(feuille_annexe.max_column)
            feuille_annexe.auto_filter.ref = (
                f"A3:{derniere_colonne}{feuille_annexe.max_row}")
    except PermissionError:
        print("❌ Ferme le fichier Excel puis relance le programme.")
        return None

    chemin = os.path.abspath(fichier)
    print(f"\n📁 Fichier créé : {chemin}")
    ouvrir_fichier(chemin)
    return chemin


# ---------------------------------------------------------------------
# Programme principal
# ---------------------------------------------------------------------

def main() -> None:
    """Enchaîne saisie, collecte, analyse et restitution."""
    lieu = demander_lieu()
    rayon_km = demander_rayon()

    stations = recuperer_stations(lieu, rayon_km)
    if stations.empty:
        print("Aucune station trouvée (ou API indisponible).")
        return

    stations = ajouter_distances(stations, lieu)
    print(f"\n✅ {len(stations)} stations trouvées à moins de "
          f"{rayon_km:g} km de {lieu.nom}")

    tous_les_prix = construire_tableau_long(stations)
    prix_fiables = tous_les_prix[tous_les_prix[COL_FIABLE]]
    prix_douteux = tous_les_prix[~tous_les_prix[COL_FIABLE]]

    classement_fiables = top3_par_carburant(
        prix_fiables, tous_les_prix, COLONNES_FIABLES,
        avec_lignes_vides=True)
    classement_douteux = top3_par_carburant(
        prix_douteux, tous_les_prix, COLONNES_DOUTEUX,
        avec_lignes_vides=False)

    afficher_classement(classement_fiables, "🏆 TOP 3 — PRIX FIABLES")
    afficher_classement(classement_douteux, "⚠️ PRIX DOUTEUX")

    if EXPORTER_EXCEL:
        exporter_excel(classement_fiables, classement_douteux, stations,
                       lieu, rayon_km)


if __name__ == "__main__":
    main()
